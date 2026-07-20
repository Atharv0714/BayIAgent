"""Generate docs/security-model.docx from the finalized security model.

Mirrors docs/security-model.html: the three-data-class model (shared internal +
per-owner private + group-protected), the permission matrix, ingest and query-path
flows, the Snowflake roles/row-policy table (owner- and group-based, with a
break-glass admin), and the locked-decisions table.

Run:  uv run --with python-docx python docs/build_security_docx.py
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Brand palette (from security-model.html) ────────────────────────
ACCENT   = RGBColor(0x7C, 0x3A, 0xED)  # purple
DEEP     = RGBColor(0xC2, 0x1E, 0x73)  # magenta/deep
OK       = RGBColor(0x1A, 0x7F, 0x52)  # green
DANGER   = RGBColor(0xE5, 0x48, 0x4D)  # red
MUTED    = RGBColor(0x6B, 0x64, 0x80)
TEXT     = RGBColor(0x1E, 0x16, 0x30)
PANEL2   = "F4F0FB"
HEADERBG = "EDE4FB"

doc = Document()

# Base font
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def heading(num, text):
    p = doc.add_paragraph()
    p.space_before = Pt(14)
    r = p.add_run(f"{num}   {text}")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = ACCENT
    return p


def subhead(text):
    p = doc.add_paragraph()
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = TEXT
    return p


def lede(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.color.rgb = MUTED
    r.font.size = Pt(10)
    return p


def note(text, kind="info"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    bar = {"info": ACCENT, "ok": OK, "warn": RGBColor(0x8A, 0x61, 0x00)}[kind]
    r = p.add_run("▍ ")
    r.font.color.rgb = bar
    r.bold = True
    r2 = p.add_run(text)
    r2.font.size = Pt(10)
    return p


def flow_steps(steps):
    """steps: list of (label, detail) rendered as a numbered decision flow."""
    for i, (label, detail) in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        rn = p.add_run(f"{i}. ")
        rn.bold = True
        rn.font.color.rgb = ACCENT
        rl = p.add_run(label)
        rl.bold = True
        if detail:
            rd = p.add_run(f" — {detail}")
            rd.font.color.rgb = MUTED


def styled_table(headers, rows, header_fill=HEADERBG, col_widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for j, h in enumerate(headers):
        set_cell_text(hdr[j], h, bold=True, size=10)
        shade(hdr[j], header_fill)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            # val may be (text, color) tuple
            if isinstance(val, tuple):
                txt, col = val
                set_cell_text(cells[j], txt, bold=(col is not None), color=col, size=10)
            else:
                set_cell_text(cells[j], val, size=10)
    if col_widths:
        for row in t.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Inches(w)
    return t


# ── Title block ─────────────────────────────────────────────────────
title = doc.add_paragraph()
tr = title.add_run("BayI")
tr.bold = True
tr.font.size = Pt(26)
tr.font.color.rgb = DEEP
tr2 = title.add_run("  Security & Access Model")
tr2.bold = True
tr2.font.size = Pt(22)
tr2.font.color.rgb = TEXT

sub = doc.add_paragraph()
sr = sub.add_run(
    "How authentication, permissions, and data sensitivity fit together once the agent runs at a "
    "shared URL. Three data classes: shared INTERNAL knowledge, PRIVATE data that only the person "
    "who ingested it can see, and PROTECTED data readable only by one privileged group."
)
sr.font.color.rgb = MUTED
sr.font.size = Pt(10.5)

status = doc.add_paragraph()
st = status.add_run("FINALIZED MODEL · DEPLOYMENT PAUSED")
st.bold = True
st.font.size = Pt(9)
st.font.color.rgb = OK

doc.add_paragraph()
lede(
    "Three enforcement layers stack from the outside in: Azure Easy Auth / SharePoint (M365) proves "
    "WHO YOU ARE (Entra, BayOne tenant only) and carries your GROUP MEMBERSHIP, the APPLICATION binds "
    "that identity + group flag to every request, and Snowflake ROW-ACCESS POLICIES are the last line "
    "that physically decides WHICH ROWS YOU CAN SEE — internal rows for everyone, private rows only "
    "for their OWNER, protected rows only for the privileged group — even if the model writes "
    "SELECT *."
)

# ── 1 · Three data classes ──────────────────────────────────────────
heading("1", "Three data classes — internal, per-owner private, group-protected")
lede(
    "Sensitivity is a property of the DATA. ANY authenticated BayOne user can flip an upload to "
    "private (owned by them, visible only to them). One privileged group can additionally mark an "
    "upload protected — visible to every member of that group, but no one else. Everything else is "
    "internal — shared knowledge every signed-in user can query."
)

styled_table(
    ["Internal — shared knowledge", "Private — owned by the ingester", "Protected — one privileged group"],
    [
        (
            "The default. Visible to every authenticated BayOne user.\n"
            "• Ingested in normal mode\n"
            "• Queryable by all staff\n"
            "• Stored with SENSITIVITY = internal\n"
            "• No owner restriction",
            "Anyone can create it. Readable only by its owner.\n"
            "• Ingested with the “Private — only me” option\n"
            "• Always ingest-only — never written to local disk\n"
            "• Stored with SENSITIVITY = private and INGESTED_BY = your email\n"
            "• Queries return your private rows — never anyone else's",
            "Only SG-BayI-Sensitive members can create or read it.\n"
            "• The “Protected — my group only” option appears only for members\n"
            "• Always ingest-only — never written to local disk\n"
            "• Stored with SENSITIVITY = protected\n"
            "• Queries return protected rows only when the caller is in the group",
        )
    ],
    col_widths=[2.17, 2.17, 2.17],
)

note(
    "Two tiers open to all, one gated. The private option is on every user's ingest screen — no "
    "membership list, isolation is by OWNER identity. The protected option appears only for members "
    "of the single privileged Entra/M365 group (SG-BayI-Sensitive); membership is decided per request "
    "from the group claim, never from row data.",
    "ok",
)
note(
    "One break-glass exception. A single named admin/compliance role (BAYI_ADMIN_READ) can read ALL "
    "private and protected data for offboarding or legal need, and every such read is logged. This is "
    "the one path by which confidential data outlives its owner (or the group) leaving the company.",
    "warn",
)

# ── 2 · Authentication ──────────────────────────────────────────────
heading("2", "Authentication — getting to the URL")
lede(
    "One HTTPS URL. Azure App Service's built-in authentication (“Easy Auth”) sits in front of "
    "the app: no valid Entra session, no bytes reach the agent. The app writes zero OAuth code — it "
    "reads the identity Azure injects."
)
subhead("Sign-in sequence")
flow_steps([
    ("User browser → Azure Easy Auth", "GET https://bayi.bayone.com over TLS."),
    ("No valid session?", "Easy Auth returns 302 redirect to the Microsoft login."),
    ("User signs in at Microsoft Entra ID", "BayOne tenant only (single-tenant)."),
    ("Entra → Easy Auth", "Returns ID token; Easy Auth sets a session cookie."),
    ("Authenticated request → app", "Easy Auth forwards it and injects the identity header "
     "X-MS-CLIENT-PRINCIPAL-NAME (email) plus the Entra/M365 group claim — the server-trusted caller "
     "identity and group membership."),
    ("App binds the caller identity + protected-group flag to the request",
     "Owner stamp on ingest; row filter on query. The app loads for the caller."),
])
note(
    "Trust boundary 1: the public internet stops at Azure Easy Auth (on the SharePoint/M365 "
    "deployment, the same Entra sign-in). Single-tenant = only BayOne Entra accounts authenticate at "
    "all. The email AND the group claim in these headers are the caller identity and membership the "
    "rest of the model relies on; the client cannot set or forge them."
)

# ── 3 · Permissions ─────────────────────────────────────────────────
heading("3", "Permissions — who can do what")
Y = ("Yes", OK)
N = ("No", DANGER)
YL = ("Yes (logged)", OK)
styled_table(
    ["Capability", "Any authenticated user\n(all BayOne staff)", "Protected-group member\n(SG-BayI-Sensitive)", "Break-glass admin\n(BAYI_ADMIN_READ · logged)"],
    [
        ["Open the app / sign in", Y, Y, Y],
        ["Query internal (shared) data", Y, Y, Y],
        ["Ingest internal documents", Y, Y, Y],
        ["Use the “Private — only me” option", Y, Y, Y],
        ["Write private rows they own (ingest-only)", Y, Y, Y],
        ["Query their OWN private rows", Y, Y, Y],
        ["Query ANOTHER user's private rows", N, N, YL],
        ["Use the “Protected — my group” option", N, Y, Y],
        ["Write protected rows (ingest-only)", N, Y, Y],
        ["Query protected group data", N, Y, YL],
    ],
    col_widths=[2.6, 1.5, 1.4, 1.5],
)
note(
    "Server-enforced, always. The owner stamped on private data and the group flag on protected data "
    "are both taken from the server-injected Easy Auth / group headers — never from client input — so "
    "a user cannot write data as someone else, cannot read rows they don't own, and cannot smuggle "
    "rows into the protected tier without being in the group. Enforcement is in the database; the UI "
    "is only convenience."
)

# ── 4 · Sensitivity selector / ingest path ──────────────────────────
heading("4", "The sensitivity selector — ingest path")
lede(
    "Every user's ingest screen offers Shared and Private; group members additionally see Protected. "
    "Choosing Private or Protected forces the INGEST-ONLY path (never written to local disk). The "
    "server stamps the owner from the authenticated identity header, and re-checks group membership "
    "before accepting a protected upload — so a non-member cannot post their way into the group tier."
)
flow_steps([
    ("Authenticated BayOne user opens Ingest", "Structure with Claude, then preview."),
    ("Choose sensitivity", "Internal / shared, Private / only me, or Protected / my group (members only)."),
    ("Internal → commit as SENSITIVITY = internal", "Draft may be saved."),
    ("Private → forced ingest-only", "Nothing kept locally; server stamps INGESTED_BY = caller."),
    ("Protected → server re-checks group membership",
     "Non-members are rejected (403); members proceed ingest-only."),
    ("Commit as SENSITIVITY = private / protected",
     "Written to Snowflake KB tables via the write-only role; INGESTED_BY = caller from the Easy Auth "
     "header, not client input."),
])
note(
    "The write role can INSERT only — no read-back, no DDL. Both KB tables carry a SENSITIVITY column "
    "(internal / private / protected) and an INGESTED_BY column (the ingester's email; empty for "
    "internal rows). Protected rows are scoped by GROUP at query time, so INGESTED_BY on them is an "
    "audit trail rather than the access key."
)

# ── 5 · Query path ──────────────────────────────────────────────────
heading("5", "Query path — how isolation is actually enforced")
lede(
    "This is the load-bearing part. Since the agent writes its own SQL, isolation can't rely on the "
    "app filtering nicely. Enforcement moves INTO Snowflake: one read role plus a ROW-ACCESS POLICY "
    "that compares each private row's INGESTED_BY to the caller and gates each protected row on the "
    "caller's group flag. The app binds both the caller identity and the protected-group flag to the "
    "Snowflake session before the agent runs its query."
)
flow_steps([
    ("User asks a question", ""),
    ("App binds BAYI_CALLER = identity and BAYI_PROTECTED = in-group? to the session",
     "From the server-trusted Easy Auth email and Entra/M365 group claim."),
    ("Run on BAYI_READ", "Admin path uses BAYI_ADMIN_READ."),
    ("Claude agent writes SQL", "SELECT-only, row cap."),
    ("Snowflake runs the query", "Row-access policy evaluates on the KB tables."),
    ("Internal row", "Returned to everyone."),
    ("Private row, INGESTED_BY = caller", "Returned to its owner."),
    ("Protected row, caller in group", "Returned to group members."),
    ("Private (other owner) / protected (non-member)", "Dropped — unless role = BAYI_ADMIN_READ."),
])
note(
    "Defense in depth: even if a prompt-injected document coaxes the model into SELECT * FROM …, a "
    "caller's connection PHYSICALLY CANNOT return another user's private rows or a non-member's "
    "protected rows — the database drops them before results leave Snowflake.",
    "warn",
)

subhead("The Snowflake pieces")
styled_table(
    ["Object", "What it does"],
    [
        ["BAYI_READ role", "SELECT on the KB tables; the row policy filters rows to internal + the "
         "caller's own private rows + protected rows when the caller's group flag is set."],
        ["BAYI_ADMIN_READ role", "Break-glass. SELECT with the row policy allowing all rows, "
         "including every user's private data and all protected data. Rare, and every use is logged."],
        ["BAYI_INGEST_WRITE role", "INSERT only, on the two ingest tables. No read, no DDL."],
        ["Row-access policy rap_ownership",
         "Returns a row when SENSITIVITY = 'internal' OR (SENSITIVITY = 'private' AND INGESTED_BY = "
         "getvariable('BAYI_CALLER')) OR (SENSITIVITY = 'protected' AND getvariable('BAYI_PROTECTED') "
         "= 'true') OR CURRENT_ROLE() = 'BAYI_ADMIN_READ'. Attached to both KB tables."],
    ],
    col_widths=[2.4, 4.1],
)
note(
    "The one trust assumption. Isolation depends on the app binding both the caller identity "
    "(BAYI_CALLER) and the protected-group flag (BAYI_PROTECTED) to the Snowflake session from the "
    "SERVER-INJECTED Easy Auth / group headers (never client input) and resetting them when a pooled "
    "connection is reused. Agent runs are already serialized on the shared connector, so setting both "
    "per request fits cleanly. The exact Snowflake mechanism (session variables read by the policy, or "
    "a per-request session) is a build-time detail to confirm."
)
note(
    "Replaces ACCOUNTADMIN. Today both read and write run as ACCOUNTADMIN (full account control). "
    "These least-privilege roles retire that for the app entirely — the single biggest risk "
    "reduction on this list.",
    "warn",
)

# ── 6 · Trust boundaries ────────────────────────────────────────────
heading("6", "Where each control lives (trust boundaries)")
flow_steps([
    ("Boundary 1 — internet → Azure", "Only authenticated BayOne accounts pass (Easy Auth gate, "
     "Entra single-tenant)."),
    ("Boundary 2 — app → Snowflake", "The app binds the caller identity and protected-group flag to "
     "the session (BAYI_READ, or BAYI_ADMIN_READ for break-glass; BAYI_INGEST_WRITE for writes); the "
     "DB enforces per-owner and per-group row visibility via the row-access policy."),
    ("Boundary 3 — app → Anthropic", "Commercial terms (no training); private and protected uploads "
     "are ingest-only and a ZDR agreement removes retention."),
])
lede(
    "Within Azure App Service: the Easy Auth gate feeds the BayI app (binds caller identity + group "
    "flag + server-side owner stamp), which persists only internal draft history to an Azure Files "
    "volume. Within the Snowflake account: BAYI_READ (plus the BAYI_ADMIN_READ break-glass role) and "
    "BAYI_INGEST_WRITE all sit behind the row-access policy that filters by owner and group flag."
)

# ── 7 · Decisions locked ────────────────────────────────────────────
heading("7", "Decisions — locked")
styled_table(
    ["#", "Decision", "Answer"],
    [
        ["1", "Sensitivity model",
         "Three data classes: internal (shared), private (per-owner), and protected (one privileged "
         "group). Reinstates the group tier above per-owner private."],
        ["2", "Who can create private data",
         "Any authenticated BayOne user — no group, no approval."],
        ["3", "Who can create protected data",
         "Only members of the single privileged Entra/M365 group (SG-BayI-Sensitive) — gated "
         "server-side at ingest."],
        ["4", "Private visibility",
         "Only the owner (the person who ingested it), enforced by the row-access policy on INGESTED_BY."],
        ["5", "Protected visibility",
         "Every member of the privileged group, enforced by the policy on a per-request group flag "
         "(BAYI_PROTECTED) — not by row owner."],
        ["6", "Break-glass",
         "One named admin/compliance role (BAYI_ADMIN_READ) can read all private and protected data, "
         "logged, for offboarding/legal."],
        ["7", "Identity & group source",
         "The Easy Auth-injected email and the Entra/M365 group claim — server-stamped, never "
         "client-settable."],
        ["8", "Private & protected uploads", "Always ingest-only — never written to local disk."],
    ],
    col_widths=[0.4, 1.7, 4.4],
)
note(
    "Three inputs still needed before build: (a) who holds the BAYI_ADMIN_READ break-glass role and "
    "how its use is logged/reviewed; (b) confirmation of the Snowflake mechanism that carries the "
    "caller identity + group flag into the session for the row-access policy to read; (c) the exact "
    "group claim SharePoint/M365 surfaces (name vs. object-id) so PROTECTED_GROUP / GROUPS_HEADER map "
    "onto it."
)

closing = doc.add_paragraph()
cr = closing.add_run(
    "Design artifact only. The build order is: least-privilege Snowflake roles (BAYI_READ, "
    "BAYI_ADMIN_READ, BAYI_INGEST_WRITE) + SENSITIVITY and INGESTED_BY columns + row policy → Easy "
    "Auth identity + group wiring → the sensitivity selector (Shared / Private / Protected, forcing "
    "ingest-only, server-stamped owner, group-gated protected) → per-request session binding of "
    "identity and group flag at query time."
)
cr.font.color.rgb = MUTED
cr.font.size = Pt(10)

out = "/Users/atharvsharma/code/BayIAgent/docs/security-model.docx"
doc.save(out)
print("Saved", out)
