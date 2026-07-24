"""Render a DocumentModel to a native .docx report (python-docx).

Title + subtitle, the headline answer, the summary paragraph, then each Table as a
real Word table with a shaded header row. python-docx has no chart primitive, so a
chartable spec is laid out as its underlying data table (labels + series) — the numbers
stay visible and auditable rather than being dropped.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from sf_agent.export.model import DocumentModel, Table, cell_text, chart_series

_ACCENT = RGBColor(0x7C, 0x3A, 0xED)
_MUTED = RGBColor(0x6B, 0x64, 0x80)
_HEADER_SHADE = "7C3AED"


def _shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): hex_color})
    tc_pr.append(shd)


def _add_table(doc: Document, table: Table) -> None:
    doc.add_heading(table.name, level=2)
    ncols = max(1, len(table.columns))
    t = doc.add_table(rows=1, cols=ncols)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, col in enumerate(table.columns):
        hdr[i].text = str(col)
        _shade_cell(hdr[i], _HEADER_SHADE)
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9)
    for row in table.rows:
        cells = t.add_row().cells
        for i in range(ncols):
            cells[i].text = cell_text(row[i]) if i < len(row) else ""
            for para in cells[i].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)


def _chart_as_table(chart: dict[str, Any]) -> Table | None:
    labels, series = chart_series(chart)
    if not labels or not series:
        return None
    columns = ["Category"] + [name for name, _ in series]
    rows = [
        [lab] + [(data[i] if i < len(data) else None) for _, data in series]
        for i, lab in enumerate(labels)
    ]
    name = str(chart.get("title") or "Chart data")
    return Table(name=name, columns=columns, rows=rows)


def render(model: DocumentModel) -> bytes:
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run(model.title)
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = _ACCENT

    sub = doc.add_paragraph()
    srun = sub.add_run(f"{model.subtitle} · {model.generated_on}")
    srun.font.size = Pt(9)
    srun.font.color.rgb = _MUTED

    if model.headline:
        hp = doc.add_paragraph()
        hlabel = hp.add_run("Answer: ")
        hlabel.font.bold = True
        hval = hp.add_run(model.headline)
        hval.font.bold = True
        hval.font.color.rgb = _ACCENT

    if model.summary:
        doc.add_paragraph(model.summary)

    if model.chart:
        ct = _chart_as_table(model.chart)
        if ct is not None:
            _add_table(doc, ct)

    for table in model.tables:
        _add_table(doc, table)

    if model.sources:
        src = doc.add_paragraph()
        srun = src.add_run("Sources: " + ", ".join(model.sources))
        srun.font.italic = True
        srun.font.size = Pt(8)
        srun.font.color.rgb = _MUTED
        src.alignment = WD_ALIGN_PARAGRAPH.LEFT

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
