"""Render a DocumentModel to a native .xlsx workbook (openpyxl).

A Summary sheet carries the headline/answer/sources; each Table becomes its own sheet
with a bold header; a chartable spec adds a Chart sheet with a native Excel chart driven
by real cells (so the chart stays live and editable, not a pasted image).
"""

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from sf_agent.export.model import DocumentModel, cell_text, chart_series

_ACCENT = "7C3AED"
_MAX_COL_WIDTH = 60


def _safe_title(name: str, used: set[str]) -> str:
    """A valid, unique Excel sheet title: strip forbidden chars, cap at 31, de-dupe."""
    clean = re.sub(r"[\[\]:*?/\\]", " ", name).strip() or "Sheet"
    clean = clean[:31]
    base = clean
    n = 2
    while clean.lower() in used:
        suffix = f" ({n})"
        clean = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(clean.lower())
    return clean


def _autosize(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        width = 0
        for cell in ws[letter]:
            v = cell.value
            if v is not None:
                width = max(width, len(str(v)))
        ws.column_dimensions[letter].width = min(max(width + 2, 10), _MAX_COL_WIDTH)


def _xl_value(v: Any) -> Any:
    """Keep numbers/bools native so Excel can sum/chart them; stringify the rest."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, bool)):
        return v
    if isinstance(v, str):
        return v
    return cell_text(v)


def _write_table_sheet(wb: Workbook, name: str, columns: list[str], rows: list[list[Any]], used: set[str]) -> Any:
    ws = wb.create_sheet(_safe_title(name, used))
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _header_fill()
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append([_xl_value(v) for v in row])
    ws.freeze_panes = "A2"
    _autosize(ws, len(columns))
    return ws


def _header_fill():
    from openpyxl.styles import PatternFill

    return PatternFill(start_color=_ACCENT, end_color=_ACCENT, fill_type="solid")


def _add_chart(wb: Workbook, chart: dict[str, Any], used: set[str]) -> None:
    labels, series = chart_series(chart)
    if not labels or not series:
        return
    ws = wb.create_sheet(_safe_title("Chart Data", used))
    ws.append(["Label"] + [name for name, _ in series])
    for i, lab in enumerate(labels):
        ws.append([lab] + [(data[i] if i < len(data) else None) for _, data in series])

    ctype = str(chart.get("type") or "bar").lower()
    if ctype in ("line", "area"):
        obj = LineChart()
    elif ctype in ("pie", "doughnut"):
        obj = PieChart()
    else:
        obj = BarChart()
        obj.type = "col"
    obj.title = chart.get("title") or None
    if hasattr(obj, "x_axis") and chart.get("x_label"):
        obj.x_axis.title = chart["x_label"]
    if hasattr(obj, "y_axis") and chart.get("y_label"):
        obj.y_axis.title = chart["y_label"]

    nrows = len(labels)
    ncols = len(series)
    data_ref = Reference(ws, min_col=2, max_col=1 + ncols, min_row=1, max_row=1 + nrows)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=1 + nrows)
    obj.add_data(data_ref, titles_from_data=True)
    obj.set_categories(cats_ref)
    if isinstance(obj, PieChart):
        # A pie shows one series; drop any extras so it renders cleanly.
        obj.series = obj.series[:1]
    obj.height = 9
    obj.width = 18
    ws.add_chart(obj, f"{get_column_letter(ncols + 3)}2")


def render(model: DocumentModel) -> bytes:
    wb = Workbook()
    used: set[str] = set()
    summary = wb.active
    summary.title = _safe_title("Summary", used)

    summary["A1"] = model.title
    summary["A1"].font = Font(bold=True, size=15, color=_ACCENT)
    summary["A2"] = f"{model.subtitle} · {model.generated_on}"
    summary["A2"].font = Font(size=10, color="6B6480")
    row = 4
    if model.headline:
        summary[f"A{row}"] = "Answer"
        summary[f"A{row}"].font = Font(bold=True, size=11)
        summary[f"B{row}"] = model.headline
        summary[f"B{row}"].font = Font(bold=True, size=11, color=_ACCENT)
        row += 1
    if model.summary:
        cell = summary[f"A{row}"]
        cell.value = model.summary
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        summary.row_dimensions[row].height = 60
        row += 2
    if model.sources:
        summary[f"A{row}"] = "Sources: " + ", ".join(model.sources)
        summary[f"A{row}"].font = Font(size=9, italic=True, color="6B6480")
    summary.column_dimensions["A"].width = 22
    for col in "BCDEF":
        summary.column_dimensions[col].width = 16

    # A deck becomes a "Slides" sheet: one row per slide (bullets joined, notes alongside).
    if model.slides:
        rows = [
            [i, s.title, "\n".join(s.bullets), s.notes or ""]
            for i, s in enumerate(model.slides, start=1)
        ]
        _write_table_sheet(wb, "Slides", ["#", "Title", "Content", "Notes"], rows, used)

    for table in model.tables:
        _write_table_sheet(wb, table.name, table.columns, table.rows, used)

    if model.chart:
        _add_chart(wb, model.chart, used)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
