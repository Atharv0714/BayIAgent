"""Local layout-preserving extraction (sf_agent.extract) and the vision-fallback routing.

The point of these tests is the distinction that costs money and correctness: a document whose
text extracts locally goes to the cheap main provider, and a POSITIONAL document must keep its
column structure on the way (otherwise an org chart silently becomes a flat name list and the
hierarchy is lost). Anything only vision can read must still route to vision.
"""

import io

from sf_agent.extract import docx_text, pdf_text, pptx_text
from sf_agent.ingest import file_content_block, local_text, needs_vision


def _org_chart_pdf() -> bytes:
    """A PDF org chart in positional form: column x encodes seniority."""
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for name, col, y in [
        ("Jane Doe", 0, 760), ("Bob Smith", 1, 730), ("Carol Lee", 2, 700),
        ("Dan Ray", 3, 670), ("Eve Ng", 3, 640), ("Heidi Vance", 1, 610),
    ]:
        c.drawString(60 + col * 120, y, name)
    c.save()
    return buf.getvalue()


def _prose_pdf() -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    lines = [
        "A rapidly scaling automotive enterprise, undergoing aggressive cost optimization,",
        "was looking to build a platform to provide resource planning for their programs.",
        "The management team lacked visibility into project workflows, leading to blind spots",
        "that compromised operational resource utilization and delayed strategic initiatives.",
    ]
    for i, line in enumerate(lines):
        c.drawString(60, 760 - i * 24, line)
    c.save()
    return buf.getvalue()


def _scanned_pdf() -> bytes:
    """A PDF with only an image and no text layer — the vision-fallback case."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    png = io.BytesIO()
    try:
        from PIL import Image

        Image.new("RGB", (60, 40), (200, 200, 200)).save(png, format="PNG")
    except ImportError:  # pragma: no cover — Pillow ships with reportlab in this project
        return b"%PDF-1.4 no text"
    png.seek(0)
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawImage(ImageReader(png), 50, 400, width=200, height=120)
    c.save()
    return buf.getvalue()


def test_positional_pdf_keeps_its_column_grid():
    """The whole point: an org chart's hierarchy lives in the column offsets, so extraction
    must emit a grid. Flat text here would read as complete while being structurally wrong."""
    text = pdf_text(_org_chart_pdf()) or ""
    assert "Jane Doe" in text
    # Depth is encoded as leading empty CSV cells, the same shape the spreadsheet lane emits.
    assert ",Bob Smith" in text
    assert ",,Carol Lee" in text
    assert ",,,Dan Ray" in text


def test_prose_pdf_is_not_shredded_into_pseudo_columns():
    """Justified prose has wide inter-word gaps that mimic column boundaries; a page of
    sentences must come back as sentences."""
    text = pdf_text(_prose_pdf()) or ""
    assert ",," not in text
    assert "automotive enterprise" in text
    assert "cost optimization" in text


def test_scanned_pdf_returns_none_so_caller_uses_vision():
    assert pdf_text(_scanned_pdf()) is None


def test_needs_vision_is_content_aware():
    # Text-based documents extract locally -> cheap main provider.
    assert needs_vision("org.pdf", _org_chart_pdf()) is False
    assert needs_vision("notes.pdf", _prose_pdf()) is False
    # No text layer, or an image -> vision.
    assert needs_vision("scan.pdf", _scanned_pdf()) is True
    assert needs_vision("chart.png", b"\x89PNG whatever") is True
    # Without bytes, stay conservative for document types.
    assert needs_vision("unknown.pdf") is True
    # Text/spreadsheet lanes never need vision.
    assert needs_vision("data.csv", b"a,b\n1,2\n") is False


def test_text_document_becomes_a_text_block_not_a_pdf_block():
    """A locally-readable PDF must be sent as text (any model can read it), not as a
    `document` block that only a vision provider can decode."""
    block = file_content_block("org.pdf", _org_chart_pdf())
    assert block["type"] == "text"
    assert ",,Carol Lee" in block["text"]


def test_scanned_document_still_becomes_a_document_block():
    block = file_content_block("scan.pdf", _scanned_pdf())
    assert block["type"] == "document"


def test_pptx_extraction_includes_speaker_notes():
    """Native deck reading beats the old PDF round-trip: a PDF render drops notes entirely."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Delivered a modernized loyalty platform for the client."
    slide.notes_slide.notes_text_frame.text = "Mention the 60% code reduction."
    buf = io.BytesIO()
    prs.save(buf)

    text = pptx_text(buf.getvalue()) or ""
    assert "modernized loyalty platform" in text
    assert "60% code reduction" in text  # notes survived
    assert local_text("deck.pptx", buf.getvalue()) is not None


def test_docx_extraction_serializes_tables():
    from docx import Document

    doc = Document()
    doc.add_paragraph("BayOne delivered a governance framework for the client this year.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Client"
    t.cell(0, 1).text = "Outcome"
    t.cell(1, 0).text = "Rivian"
    t.cell(1, 1).text = "60% code reduction"
    buf = io.BytesIO()
    doc.save(buf)

    text = docx_text(buf.getvalue()) or ""
    assert "governance framework" in text
    assert "| Client | Outcome |" in text
    assert "Rivian" in text
