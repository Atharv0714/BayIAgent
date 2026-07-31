"""Template-based pptx rendering: build the deck onto an uploaded template's layouts."""

import io

from pptx import Presentation

from sf_agent.export import answer_to_document, render_document
from sf_agent.export.render_pptx import _SLIDE_W

_DECK = {
    "answer": "Retail capabilities deck.",
    "values": {
        "slides": [
            {"title": "Retail Wins", "content": ["Albertsons UI", "Walmart marketing"],
             "speaker_notes": "lead with the wins"},
            {"title": "The Ask", "content": ["Partner with BayOne"]},
        ]
    },
}


def _template_bytes() -> bytes:
    """A stand-in template: python-pptx's built-in deck (has Title + Title-and-Content
    layouts, and its own slide size distinct from our widescreen default)."""
    buf = io.BytesIO()
    Presentation().save(buf)
    return buf.getvalue()


def test_template_deck_inherits_template_size_and_populates_slides():
    tpl = _template_bytes()
    tpl_width = Presentation(io.BytesIO(tpl)).slide_width  # the template's own size

    model = answer_to_document(_DECK, "Costco retail deck")
    data = render_document(model, "pptx", template=tpl)

    prs = Presentation(io.BytesIO(data))
    assert prs.slide_width == tpl_width  # template's size governs, not our widescreen override
    slides = list(prs.slides)
    assert len(slides) == 3  # title + 2 content slides (the template's sample slides removed)

    # First content slide: title + bullets landed in the layout's placeholders.
    texts = [sh.text for sh in slides[1].shapes if sh.has_text_frame]
    assert any("Retail Wins" in t for t in texts)
    assert any("Albertsons UI" in t for t in texts)
    # Speaker notes carried onto the notes page.
    assert slides[1].has_notes_slide
    assert "lead with the wins" in slides[1].notes_slide.notes_text_frame.text


def test_default_path_unchanged_without_template():
    model = answer_to_document(_DECK, "deck")
    data = render_document(model, "pptx")  # no template
    prs = Presentation(io.BytesIO(data))
    assert prs.slide_width == int(_SLIDE_W)  # our widescreen default
    assert len(list(prs.slides)) >= 3


def test_template_render_reopens_and_is_valid():
    model = answer_to_document(_DECK, "deck")
    data = render_document(model, "pptx", template=_template_bytes())
    assert data[:2] == b"PK"
    # re-open must succeed
    assert len(list(Presentation(io.BytesIO(data)).slides)) == 3


# --- templates whose design lives in their slides, not their theme ---------------------
# A corporate .potx carries its identity in the master/layouts/theme, so inheriting those
# reproduces the look. Decks exported from Canva/Figma/Slides — and any deck built by a
# script — ship the stock Office theme and one EMPTY layout, drawing every background,
# colour and font directly on the slide. Keeping the theme and dropping the slides then
# keeps nothing, and the user gets white slides in default Calibri: "a blank deck".


def _designed_template() -> bytes:
    """A template in that shape: empty layout, look carried entirely by the slides."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    for heading, body in (("Cover", "the opening copy"), ("Detail", "a good deal more body copy here")):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank: no placeholders
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0xF5, 0xF5, 0xF5)

        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1)).text_frame
        run = title.paragraphs[0].add_run()
        run.text = heading
        run.font.name, run.font.size, run.font.bold = "Figtree", Pt(36), True
        run.font.color.rgb = RGBColor(0x17, 0x17, 0x17)

        copy = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(8), Inches(3)).text_frame
        crun = copy.paragraphs[0].add_run()
        crun.text = body
        crun.font.name, crun.font.size = "Nunito Sans", Pt(12)
        crun.font.color.rgb = RGBColor(0x52, 0x52, 0x52)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_look_is_read_off_the_slides_when_layouts_are_empty():
    from sf_agent.export.render_pptx import _extract_look
    from pptx import Presentation

    look = _extract_look(Presentation(io.BytesIO(_designed_template())))
    assert look.background is not None and b"F5F5F5" in look.background
    assert look.title.name == "Figtree" and look.title.bold is True
    assert str(look.title.color) == "171717"
    assert look.body.name == "Nunito Sans"
    assert str(look.body.color) == "525252"
    # "not bold" is a real answer — it must not be back-filled from the bold titles.
    assert look.body.bold is not True


def test_deck_on_a_designed_template_is_not_blank():
    from pptx import Presentation

    from sf_agent.export import answer_to_document, render_document

    model = answer_to_document(
        {"answer": "x", "values": {"slides": [
            {"title": "Generated One", "content": ["first point", "second point"]},
            {"title": "Generated Two", "content": ["third point"]},
        ]}},
        "deck",
    )
    data = render_document(model, "pptx", template=_designed_template(), template_mode="theme")
    prs = Presentation(io.BytesIO(data))

    for slide in prs.slides:
        assert slide._element.xpath("./p:cSld/p:bg"), "template background not carried over"

    fonts, colors = set(), set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    fonts.add(run.font.name)
                    if run.font.color is not None and run.font.color.type is not None:
                        colors.add(str(run.font.color.rgb))
    assert "Figtree" in fonts and "Nunito Sans" in fonts, f"template fonts missing: {fonts}"
    assert {"171717", "525252"} <= colors, f"template colours missing: {colors}"
    assert None not in fonts, "some text was left on the default theme font"


def test_a_template_with_real_placeholders_is_left_alone():
    """The look is only read when there is nothing to inherit; a proper template must keep
    filling its own placeholders and must not get a slide-level background forced onto it."""
    from pptx import Presentation

    from sf_agent.export import answer_to_document, render_document

    tpl = Presentation()
    slide = tpl.slides.add_slide(tpl.slide_layouts[1])
    slide.shapes.title.text = "Cover"
    slide.placeholders[1].text_frame.text = "template body"
    buf = io.BytesIO()
    tpl.save(buf)

    model = answer_to_document({"answer": "x", "values": {"slides": [
        {"title": "Generated", "content": ["a"]}]}}, "deck")
    prs = Presentation(io.BytesIO(render_document(model, "pptx", template=buf.getvalue())))
    for s in prs.slides:
        assert not s._element.xpath("./p:cSld/p:bg"), "background forced onto a themed template"
        assert any(sh.is_placeholder for sh in s.shapes), "stopped using the template's placeholders"
