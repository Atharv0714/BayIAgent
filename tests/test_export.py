"""Export module: answer -> document mapping and each renderer producing a valid file."""

import datetime as dt
import io

import pytest

from sf_agent.export import (
    SUPPORTED_FORMATS,
    answer_to_document,
    humanize,
    render_document,
)
from sf_agent.export.model import Table, cell_text, chart_series

# A representative /api/ask answer: a headline value, a list-of-objects table (with a
# null cell to prove gaps survive), a scalar list, a bare scalar, and a chart.
ANSWER = {
    "answer": "There are 127 distinct clients; the top 3 by placements are shown.",
    "value": 127,
    "values": {
        "top_clients": [
            {"client": "HPE", "placements": 42, "end_date": None},
            {"client": "Cisco", "placements": 31, "end_date": "2025-01-01"},
            {"client": "Oracle", "placements": 20, "end_date": None},
        ],
        "skills": ["Python", "SQL", "React"],
        "note": "some placements have no end date",
    },
    "chart": {
        "type": "bar", "title": "Top clients", "x_label": "Client", "y_label": "Placements",
        "labels": ["HPE", "Cisco", "Oracle"], "series": [{"name": "Placements", "data": [42, 31, 20]}],
    },
    "sources": ["BILLABLE_DATA_BY_JOB_DESCRIPTION"],
}

_MAGIC = {"xlsx": b"PK", "docx": b"PK", "pptx": b"PK", "pdf": b"%PDF"}


def test_humanize_keeps_acronyms_and_titlecases() -> None:
    assert humanize("candidate_full_name") == "Candidate Full Name"
    assert humanize("client_id") == "Client ID"
    assert humanize("SOW-number") == "SOW Number"


def test_cell_text_flags_missing_and_serializes() -> None:
    assert cell_text(None) == "—"
    assert cell_text("") == "—"
    assert cell_text(True) == "Yes"
    assert cell_text({"a": 1}) == '{"a": 1}'


def test_answer_to_document_maps_shapes() -> None:
    model = answer_to_document(ANSWER, "How many clients?")
    assert model.title == "How many clients?"
    assert model.headline == "127"
    names = {t.name: t for t in model.tables}
    # list-of-objects -> real table with humanized headers and the null preserved
    top = names["Top Clients"]
    assert top.columns == ["Client", "Placements", "End Date"]
    assert top.rows[0] == ["HPE", 42, None]
    # list-of-scalars -> single-column table
    assert names["Skills"].rows == [["Python"], ["SQL"], ["React"]]
    # bare scalar -> folded into a Details key/value table
    assert ["Note", "some placements have no end date"] in names["Details"].rows
    assert model.chart is not None


def test_answer_to_document_defaults_title_when_no_question() -> None:
    model = answer_to_document({"answer": "x", "value": None}, "")
    assert model.title == "BayI Intelligence Report"
    assert model.headline is None


def test_chart_series_drops_nonnumeric() -> None:
    labels, series = chart_series(ANSWER["chart"])
    assert labels == ["HPE", "Cisco", "Oracle"]
    assert series == [("Placements", [42.0, 31.0, 20.0])]
    # a scatter-style spec (dict data) yields no chartable series
    _, empty = chart_series({"labels": ["a"], "series": [{"name": "s", "data": [{"x": 1, "y": 2}]}]})
    assert empty == []


@pytest.mark.parametrize("fmt", SUPPORTED_FORMATS)
def test_renderer_produces_valid_file(fmt: str) -> None:
    model = answer_to_document(ANSWER, "How many clients?")
    data = render_document(model, fmt)
    assert isinstance(data, bytes) and len(data) > 100
    assert data[: len(_MAGIC[fmt])] == _MAGIC[fmt]


@pytest.mark.parametrize("fmt", SUPPORTED_FORMATS)
def test_renderer_handles_empty_answer(fmt: str) -> None:
    # A minimal answer (no values, no chart) must still render a file, not crash.
    model = answer_to_document({"answer": "No data found.", "value": None}, "q")
    data = render_document(model, fmt)
    assert data[: len(_MAGIC[fmt])] == _MAGIC[fmt]


def test_office_files_reopen() -> None:
    model = answer_to_document(ANSWER, "How many clients?")
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(render_document(model, "xlsx")))
    assert "Summary" in wb.sheetnames and "Top Clients" in wb.sheetnames

    from docx import Document

    doc = Document(io.BytesIO(render_document(model, "docx")))
    assert len(doc.tables) >= 1

    from pptx import Presentation

    prs = Presentation(io.BytesIO(render_document(model, "pptx")))
    assert len(list(prs.slides)) >= 2


def test_unsupported_format_raises() -> None:
    model = answer_to_document(ANSWER, "q")
    with pytest.raises(ValueError):
        render_document(model, "rtf")


# --- deck generation (values.slides -> real slides) ------------------------------------
_DECK = {
    "answer": "Here is the 3-slide deck. It is ready to download.",
    "value": None,
    "values": {
        "slides": [
            {"title": "Overview", "content": ["Point A", "Point B"], "speaker_notes": "intro"},
            {"title": "Market", "content": ["TAM is large"]},
            {"title": "Ask", "content": ["Partner with BayOne"], "speaker_notes": "close"},
        ]
    },
}


def test_slides_parse_into_deck_not_table():
    m = answer_to_document(_DECK, "Make a deck")
    assert len(m.slides) == 3
    assert m.tables == []  # the slides array is a deck, not a generic table
    assert m.slides[0].title == "Overview"
    assert m.slides[0].bullets == ["Point A", "Point B"]
    assert m.slides[0].notes == "intro"
    assert m.slides[1].notes is None


def test_slides_content_as_string_splits_to_bullets():
    m = answer_to_document(
        {"answer": "x", "values": {"slides": [{"title": "T", "content": "line one\nline two"}]}}, "q"
    )
    assert m.slides[0].bullets == ["line one", "line two"]


def test_pptx_emits_one_slide_per_entry():
    import io as _io

    from pptx import Presentation

    prs = Presentation(_io.BytesIO(render_document(answer_to_document(_DECK, "Make a deck"), "pptx")))
    slides = list(prs.slides)
    # title + summary + 3 content slides
    assert len(slides) >= 4
    # speaker notes carried onto the notes page
    assert any(s.has_notes_slide and "intro" in s.notes_slide.notes_text_frame.text for s in slides)


def test_slides_named_column_is_not_treated_as_a_deck():
    # A numeric 'slides' value (not a list of slide objects) stays an ordinary value.
    m = answer_to_document({"answer": "x", "value": 5, "values": {"slides": 5}}, "how many slides?")
    assert m.slides == []


# --- markdown must not leak into generated documents -----------------------------------
def test_strip_markdown_removes_markers_but_keeps_text():
    from sf_agent.export.model import strip_markdown

    assert strip_markdown("**58 clients** via `MASTERSKILLLIST`") == "58 clients via MASTERSKILLLIST"
    assert strip_markdown("### Heading") == "Heading"
    assert strip_markdown("- **Google** (145)") == "Google (145)"
    assert strip_markdown("*eBay* and _Cisco_") == "eBay and Cisco"
    # Not markdown: arithmetic and snake_case identifiers must survive untouched.
    assert strip_markdown("Rate is 5 * 3 and a_var_name stays") == "Rate is 5 * 3 and a_var_name stays"


def test_document_summary_and_slides_are_plain_text():
    """The prompt asks for markdown, the chat UI renders it — but PPTX/DOCX/PDF writers emit
    plain runs, so markers would show as literal '**' in a client-facing deck."""
    model = answer_to_document(
        {
            "answer": "We have **650** billable resources — **418** Enterprise.",
            "values": {
                "slides": [
                    {
                        "title": "**Overview**",
                        "content": ["- **Google** 145", "*eBay* 38"],
                        "speaker_notes": "Cite **exactly**.",
                    }
                ]
            },
        },
        "q",
    )
    assert "**" not in model.summary
    assert model.summary == "We have 650 billable resources — 418 Enterprise."
    s = model.slides[0]
    assert s.title == "Overview"
    assert s.bullets == ["Google 145", "eBay 38"]
    assert s.notes == "Cite exactly."
    # And it survives into a real file.
    assert render_document(model, "pptx")[:2] == b"PK"


# --- consistent formatting reaches the documents ----------------------------------------
_FMT_ANSWER = {
    "answer": "Top clients by billing run-rate.",
    "value": 13278.7,
    "values": {
        "top_clients": [
            {"client_name": "Google", "run_rate_usd": 13278.7, "gm": 12.46,
             "placement_count": 145, "end_date": "2026-07-31", "companyid": 1234567},
            {"client_name": "Cisco", "run_rate_usd": 3439.74, "gm": -89.5,
             "placement_count": 85, "end_date": "2026-12-31", "companyid": 7654321},
        ],
        "annual_salary": 140000,
    },
    "sources": ["V_CLIENT_PLACEMENTS"],
}


def test_docx_figures_match_the_screen_formatting():
    from docx import Document

    model = answer_to_document(_FMT_ANSWER, "Top clients by revenue")
    doc = Document(io.BytesIO(render_document(model, "docx")))
    text = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "$13,278.70" in text          # currency grouped, 2dp
    assert "12.46%" in text and "-89.5%" in text   # percent, negatives intact
    assert "Jul 31, 2026" in text        # date humanised
    assert "2026-07-31" not in text      # ...and the raw ISO form is gone
    assert "1234567" in text and "1,234,567" not in text  # identifiers never grouped
    assert "/hr" not in text             # no invented rate unit


def test_xlsx_keeps_values_native_so_excel_can_compute():
    """The workbook must stay numeric — presentation goes in the cell number format."""
    from openpyxl import load_workbook

    model = answer_to_document(_FMT_ANSWER, "Top clients by revenue")
    ws = load_workbook(io.BytesIO(render_document(model, "xlsx")))["Top Clients"]
    header = [c.value for c in ws[1]]
    cells = {name: ws.cell(2, header.index(name) + 1) for name in header}

    assert isinstance(cells["Run Rate USD"].value, float)
    assert cells["Run Rate USD"].number_format == '"$"#,##0.00'
    assert isinstance(cells["Placement Count"].value, int)
    # An ISO date becomes a real date, or Excel cannot sort/filter/chart it.
    assert isinstance(cells["End Date"].value, (dt.date, dt.datetime))
    assert cells["End Date"].number_format == "mmm d, yyyy"
    assert cells["Companyid"].number_format == "0"
