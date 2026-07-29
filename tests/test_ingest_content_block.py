"""Offline tests for the ingest content-block builder (sf_agent.ingest)."""

import base64

import pytest

from sf_agent import ingest
from sf_agent.ingest import IngestError, build_content_block


def test_pdf_becomes_a_document_block():
    raw = b"%PDF-1.4 fake"
    blocks = build_content_block("deck.pdf", raw)
    assert blocks[0]["type"] == "document"
    src = blocks[0]["source"]
    assert src["media_type"] == "application/pdf"
    assert base64.b64decode(src["data"]) == raw
    # A trailing instruction text block is always appended.
    assert blocks[-1]["type"] == "text"
    assert "deck.pdf" in blocks[-1]["text"]


def test_png_becomes_an_image_block():
    blocks = build_content_block("chart.PNG", b"\x89PNG data")
    assert blocks[0]["type"] == "image"
    assert blocks[0]["source"]["media_type"] == "image/png"


def test_jpg_and_jpeg_map_to_jpeg_media_type():
    for name in ("photo.jpg", "photo.jpeg"):
        blocks = build_content_block(name, b"\xff\xd8\xff data")
        assert blocks[0]["source"]["media_type"] == "image/jpeg"


def test_text_family_is_decoded_inline_with_filename():
    blocks = build_content_block("data.csv", b"client,revenue\nHPE,1800000\n")
    assert blocks[0]["type"] == "text"
    assert "Filename: data.csv" in blocks[0]["text"]
    assert "HPE,1800000" in blocks[0]["text"]


def test_text_decode_is_lossy_not_fatal():
    # Invalid UTF-8 must not raise — errors="replace" keeps ingestion going.
    blocks = build_content_block("notes.txt", b"caf\xe9 \xff\xfe")
    assert blocks[0]["type"] == "text"
    assert "caf" in blocks[0]["text"]


def test_office_file_is_converted_to_a_pdf_block(monkeypatch):
    # LibreOffice is stubbed so the test stays offline; the office branch must route the
    # converted bytes through the PDF path (document block, application/pdf).
    monkeypatch.setattr(ingest, "_office_to_pdf", lambda filename, raw: b"%PDF-1.7 converted")
    blocks = build_content_block("deck.pptx", b"PK\x03\x04")
    assert blocks[0]["type"] == "document"
    assert blocks[0]["source"]["media_type"] == "application/pdf"
    assert base64.b64decode(blocks[0]["source"]["data"]) == b"%PDF-1.7 converted"
    # Provenance keeps the original office filename, not a .pdf rename.
    assert "deck.pptx" in blocks[-1]["text"]


def test_office_without_libreoffice_raises_convert_message(monkeypatch):
    monkeypatch.setattr(ingest, "_find_soffice", lambda: None)
    with pytest.raises(IngestError) as exc:
        build_content_block("report.docx", b"PK\x03\x04")
    assert "LibreOffice is not installed" in str(exc.value)


def test_unknown_extension_is_rejected():
    with pytest.raises(IngestError):
        build_content_block("archive.zip", b"PK\x03\x04")


def _tiny_xlsx() -> bytes:
    """A 2-sheet workbook with a positional (org-chart-like) grid on sheet 1."""
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Org 2026"
    ws["A1"] = "Jane Doe"
    ws["B2"] = "Bob Smith"      # one column right, one row down: Bob reports to Jane
    ws["C3"] = "Carol Lee"
    ws2 = wb.create_sheet("VP")
    ws2["A1"] = "VPs with Customer"
    ws2["A2"] = "Pvv Raju"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_becomes_positional_csv_text_not_pdf():
    """Excel must reach the model as per-sheet CSV TEXT preserving the grid — never via the
    LibreOffice->PDF render, whose pagination scatters a wide sheet's columns across pages
    and silently destroys positional documents like org charts."""
    blocks = build_content_block("Customer_Mapping.xlsx", _tiny_xlsx())
    assert blocks[0]["type"] == "text"  # not a 'document' (PDF) block
    text = blocks[0]["text"]
    assert "Filename: Customer_Mapping.xlsx" in text
    # Both sheets present, labeled.
    assert "### Sheet: Org 2026" in text and "### Sheet: VP" in text
    # The grid POSITIONS survive: Bob is one column right of Jane, Carol two right.
    assert "Jane Doe" in text
    assert ",Bob Smith" in text
    assert ",,Carol Lee" in text


def test_xlsx_never_calls_libreoffice(monkeypatch):
    """Guard the routing itself: even with LibreOffice present, xlsx must not go to PDF."""

    def _boom(*a, **k):  # pragma: no cover - failing is the assertion
        raise AssertionError("xlsx was routed through _office_to_pdf")

    monkeypatch.setattr(ingest, "_office_to_pdf", _boom)
    blocks = build_content_block("grid.xlsx", _tiny_xlsx())
    assert blocks[0]["type"] == "text"


def test_xls_still_uses_office_path(monkeypatch):
    """Legacy .xls (openpyxl can't read it) must still take the LibreOffice route."""
    called = {}

    def _fake(filename, raw):
        called["yes"] = True
        return b"%PDF-1.4 converted"

    monkeypatch.setattr(ingest, "_office_to_pdf", _fake)
    blocks = build_content_block("old.xls", b"legacy bytes")
    assert called.get("yes") and blocks[0]["type"] == "document"


def test_corrupt_xlsx_raises_ingest_error():
    with pytest.raises(IngestError, match="Excel"):
        build_content_block("broken.xlsx", b"not a zip at all")
