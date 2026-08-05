"""Chat file upload: reusable content-block builder, /api/upload classify+store, and
export accepting a template_id. No live warehouse or model needed."""

import asyncio
import io
import json
import posixpath

import pytest
from fastapi import Request, UploadFile

from sf_agent.ingest import IngestError, file_content_block
from sf_agent.web import ExportRequest, STATE, export, upload


@pytest.fixture(autouse=True)
def _clear_uploads():
    STATE.uploads.clear()
    yield
    STATE.uploads.clear()


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    buf = io.BytesIO()
    Presentation().save(buf)
    return buf.getvalue()


def _request(identity: str | None = None):
    """A bare ASGI request, optionally carrying an Easy Auth identity header."""
    headers = [(b"x-ms-client-principal-name", identity.encode())] if identity else []
    return Request({"type": "http", "method": "POST", "path": "/", "headers": headers})


def _upload(filename: str, data: bytes, identity: str | None = None):
    return asyncio.run(
        upload(_request(identity), UploadFile(file=io.BytesIO(data), filename=filename))
    )


def _body(resp) -> dict:
    return json.loads(bytes(resp.body))


# --- file_content_block (reused by ingest + chat context) ------------------------------

def test_file_content_block_text_returns_text_block():
    block = file_content_block("notes.txt", b"hello world")
    assert block["type"] == "text"
    assert "hello world" in block["text"]


def test_file_content_block_unsupported_raises():
    with pytest.raises(IngestError, match="Unsupported file type"):
        file_content_block("mystery.xyz", b"\x00\x01")


# --- /api/upload classification + storage ---------------------------------------------

def test_upload_pptx_is_template_and_keeps_raw():
    resp = _upload("brand.pptx", _pptx_bytes())
    body = _body(resp)
    assert body["ok"] and body["kind"] == "template"
    rec = STATE.uploads[body["upload_id"]]
    # A deck keeps BOTH roles: raw bytes for rendering onto its theme, and a content block so
    # the model can also read it. Storing only the template made "summarize this deck"
    # impossible — the file being asked about was the one file the model could not see.
    assert rec["kind"] == "template" and isinstance(rec["raw"], bytes)
    assert "block" in rec, "a deck must also be readable as context"


def test_upload_text_is_context_and_stores_block():
    resp = _upload("brief.txt", b"BayOne retail brief")
    body = _body(resp)
    assert body["ok"] and body["kind"] == "context"
    rec = STATE.uploads[body["upload_id"]]
    assert rec["kind"] == "context" and rec["block"]["type"] == "text" and "raw" not in rec


def test_upload_reports_capabilities_not_one_fixed_role():
    """The response must say what a file CAN be used for, not commit it to a role.

    The UI labelled every .pptx "Template" because /api/upload returned only the internal
    ``kind``. The role is decided per question (describe_upload_use), so the response carries
    capability flags and the chip shows a neutral "Deck"/"Document".
    """
    deck = _body(_upload("brand.pptx", _pptx_bytes()))
    assert deck["can_template"] is True and deck["readable"] is True

    doc = _body(_upload("brief.txt", b"BayOne retail brief"))
    assert doc["can_template"] is False and doc["readable"] is True


def test_upload_empty_file_400():
    resp = _upload("empty.txt", b"")
    assert resp.status_code == 400


def test_upload_unsupported_type_415():
    resp = _upload("weird.xyz", b"\x00\x01\x02")
    assert resp.status_code == 415


# --- export accepts a template_id ------------------------------------------------------

def test_export_pptx_with_template_id_renders():
    STATE.uploads["tpl1"] = {"filename": "brand.pptx", "kind": "template", "raw": _pptx_bytes()}
    req = ExportRequest(
        format="pptx",
        question="Make a deck",
        answer={"answer": "x", "values": {"slides": [{"title": "A", "content": ["b"]}]}},
        template_id="tpl1",
    )
    resp = export(req, _request())
    assert resp.status_code == 200
    assert bytes(resp.body)[:2] == b"PK"  # a valid pptx/zip


def test_export_missing_template_id_falls_back():
    # A stale/unknown template_id must not error — it falls back to the default deck.
    req = ExportRequest(
        format="pptx",
        question="Make a deck",
        answer={"answer": "x", "values": {"slides": [{"title": "A", "content": ["b"]}]}},
        template_id="does-not-exist",
    )
    resp = export(req, _request())
    assert resp.status_code == 200
    assert bytes(resp.body)[:2] == b"PK"


# --- an attachment is used per the QUESTION, not per its extension ----------------------
def test_deck_attachment_is_read_when_the_question_asks_about_its_contents():
    """The original bug: every .pptx was a template, so "summarize this deck" could never
    work — the file being asked about was the one file the model could not see."""
    from sf_agent.export.intent import describe_upload_use

    use = describe_upload_use("Summarize this deck for me", is_deck=True)
    assert use["read"] is True and use["as_template"] is False


def test_deck_attachment_is_a_template_when_the_question_asks_for_styling():
    from sf_agent.export.intent import describe_upload_use

    for q in ("Make a deck in this template", "Build a slide deck using this theme"):
        use = describe_upload_use(q, is_deck=True)
        assert use["as_template"] is True, q


def test_deck_attachment_can_serve_both_roles_in_one_turn():
    from sf_agent.export.intent import describe_upload_use

    use = describe_upload_use("Summarize this deck and make a new one in the same theme", True)
    assert use["read"] is True and use["as_template"] is True


def test_bare_deck_attachment_defaults_to_template_use():
    from sf_agent.export.intent import describe_upload_use

    use = describe_upload_use("Our top clients by revenue", is_deck=True)
    assert use["as_template"] is True


def test_non_deck_attachment_is_always_read():
    from sf_agent.export.intent import describe_upload_use

    use = describe_upload_use("Make a deck in this theme", is_deck=False)
    assert use["read"] is True and use["as_template"] is False


def test_template_mode_distinguishes_headers_from_theme_only():
    from sf_agent.export.intent import THEME_AND_HEADERS, THEME_ONLY, detect_template_mode

    assert detect_template_mode("same headers and theme as the slide") == THEME_AND_HEADERS
    assert detect_template_mode("keep the section titles from the template") == THEME_AND_HEADERS
    assert detect_template_mode("use the same structure as the deck") == THEME_AND_HEADERS
    # A bare theme request, and an explicit "just the theme", stay theme-only.
    assert detect_template_mode("build it in the same theme") == THEME_ONLY
    assert detect_template_mode("just the theme please") == THEME_ONLY
    assert detect_template_mode("only the styling, not the headings") == THEME_ONLY
    assert detect_template_mode("") == THEME_ONLY


def test_theme_and_headers_reuses_the_templates_slide_titles():
    """theme_and_headers must carry the uploaded deck's own headings onto the new slides."""
    import io

    from pptx import Presentation
    from pptx.util import Inches

    from sf_agent.export import answer_to_document, render_document

    # A template whose slides carry distinctive headers.
    tpl = Presentation()
    for heading in ("Client Overview", "Our Approach"):
        slide = tpl.slides.add_slide(tpl.slide_layouts[5])
        slide.shapes.title.text = heading
    buf = io.BytesIO()
    tpl.save(buf)
    template = buf.getvalue()

    model = answer_to_document(
        {"answer": "x", "values": {"slides": [
            {"title": "Generated One", "content": ["a"]},
            {"title": "Generated Two", "content": ["b"]},
        ]}},
        "deck",
    )

    themed = render_document(model, "pptx", template=template, template_mode="theme")
    headed = render_document(model, "pptx", template=template, template_mode="theme_and_headers")

    def titles(data: bytes) -> str:
        prs = Presentation(io.BytesIO(data))
        return " ".join(
            sh.text_frame.text for s in prs.slides for sh in s.shapes
            if getattr(sh, "has_text_frame", False)
        )

    # theme-only keeps OUR titles; theme_and_headers adopts the TEMPLATE's.
    assert "Generated One" in titles(themed)
    assert "Client Overview" in titles(headed)


def test_deck_on_template_is_a_valid_package_powerpoint_will_open():
    """A deck rendered onto a template must not carry the template's discarded slide parts.

    Clearing only the ``<p:sldId>`` entries left the presentation still related to the old
    slide parts, so they were written to the .pptx — while python-pptx handed the NEW slides
    the same partnames (``slide{len(sldIdLst) + 1}.xml`` off a now-empty list). The result was
    a package with two ``ppt/slides/slide1.xml`` entries and slides claimed by two notes pages
    at once, which PowerPoint reports as "found a problem with content ... attempt to repair".
    """
    import collections
    import re
    import zipfile
    from xml.etree import ElementTree as ET

    from pptx import Presentation

    from sf_agent.export import answer_to_document, render_document

    # A template that looks like a real one: several slides, each with a speaker-notes page.
    tpl = Presentation()
    for heading in ("Cover", "Agenda", "Approach", "Pricing"):
        slide = tpl.slides.add_slide(tpl.slide_layouts[5])
        slide.shapes.title.text = heading
        slide.notes_slide.notes_text_frame.text = f"template note for {heading}"
    buf = io.BytesIO()
    tpl.save(buf)
    template = buf.getvalue()

    model = answer_to_document(
        {"answer": "x", "values": {"slides": [
            {"title": f"Generated {i}", "content": [f"point {i}"], "speaker_notes": f"note {i}"}
            for i in range(1, 6)
        ]}},
        "deck",
    )
    data = render_document(model, "pptx", template=template, template_mode="theme")

    z = zipfile.ZipFile(io.BytesIO(data))
    entries = z.namelist()
    names = set(entries)

    dupes = [n for n, c in collections.Counter(entries).items() if c > 1]
    assert not dupes, f"package writes the same part twice: {dupes}"

    # Every notes page must be claimed by exactly one slide, and that slide must claim it back.
    claims: dict[str, list[str]] = collections.defaultdict(list)
    for note in (n for n in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)):
        rels = f"ppt/notesSlides/_rels/{note.split('/')[-1]}.rels"
        if rels not in names:
            continue
        for rel in ET.fromstring(z.read(rels)):
            target = rel.get("Target") or ""
            if "/slide" in (rel.get("Type") or "") and "notesSlide" not in target:
                claims[target.split("/")[-1]].append(note.split("/")[-1])
    doubled = {s: n for s, n in claims.items() if len(n) > 1}
    assert not doubled, f"slides carrying more than one notes page: {doubled}"

    # No relationship may point at a part that isn't in the package.
    for rels in (n for n in names if n.endswith(".rels")):
        base = posixpath.dirname(posixpath.dirname(rels))
        for rel in ET.fromstring(z.read(rels)):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target")
            resolved = (
                target[1:] if target.startswith("/")
                else posixpath.normpath(posixpath.join(base, target))
            )
            assert resolved in names, f"{rels} points at missing part {target}"

    # ...and the deck still reads correctly: our slides, our notes, the template's theme.
    prs = Presentation(io.BytesIO(data))
    text = " ".join(
        sh.text_frame.text for s in prs.slides for sh in s.shapes
        if getattr(sh, "has_text_frame", False)
    )
    assert "Generated 1" in text and "Generated 5" in text
    assert "template note for Cover" not in text  # the template's own slides are gone
